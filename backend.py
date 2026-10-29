
import io
import re


def read_resume(uploaded_file):
    """
    Read text from an uploaded PDF or DOCX resume.
    Does not depend on the current file pointer position.
    """

    if uploaded_file is None:
        raise ValueError("No resume file was uploaded.")

    filename = uploaded_file.name.lower()

    # Read the complete uploaded file as bytes.
    file_data = uploaded_file.getvalue()

    if not file_data:
        raise ValueError("The uploaded file is empty.")

    # Read PDF files.
    if filename.endswith(".pdf"):
        from pypdf import PdfReader

        try:
            pdf_file = io.BytesIO(file_data)
            reader = PdfReader(pdf_file)

            if reader.is_encrypted:
                raise ValueError(
                    "The PDF is password-protected. "
                    "Please upload an unlocked PDF."
                )

            pages_text = []

            for page in reader.pages:
                pages_text.append(page.extract_text() or "")

            text = "\n".join(pages_text)

        except ValueError:
            raise
        except Exception as error:
            raise ValueError(
                "Unable to read this PDF. "
                "Please check that the file is valid."
            ) from error

    # Read DOCX files.
    elif filename.endswith(".docx"):
        from docx import Document

        try:
            document = Document(io.BytesIO(file_data))

            paragraphs = [
                paragraph.text
                for paragraph in document.paragraphs
            ]

            # Include text inside tables.
            for table in document.tables:
                for row in table.rows:
                    for cell in row.cells:
                        paragraphs.append(cell.text)

            text = "\n".join(paragraphs)

        except Exception as error:
            raise ValueError(
                "Unable to read this DOCX file. "
                "Please check that the file is valid."
            ) from error

    else:
        raise ValueError(
            "Unsupported file type. Upload a PDF or DOCX file."
        )

    # Clean unnecessary whitespace.
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n[ \t]*\n+", "\n", text)
    text = text.strip()

    if not text:
        raise ValueError(
            "No readable text was found. "
            "If your PDF contains scanned images, OCR is required."
        )

    return text
  
